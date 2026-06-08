"""
跨会话长期记忆测试 · 定时执行器
========================================
专门用于管理"需要等待时间间隔"的测试用例（类型B：C2/C4/C5）。
通过状态文件驱动，每次运行只推进"该执行"的session。
配合 WorkBuddy 每日自动化使用。

用法：
  # 初始化：从 timed xlsx 创建状态文件，执行所有 S1
  python3 memory_test_runner.py --init test_cases_timed.xlsx

  # 增量推进：检查状态文件，执行所有到期的下一个 session
  python3 memory_test_runner.py --step

  # 查看当前进度
  python3 memory_test_runner.py --status

  # 强制重跑某个用例的某个 session
  python3 memory_test_runner.py --retry C2-01 --session S2

状态文件: cross_session_state.json（自动管理，不要手动修改）
"""

import json
import os
import sys
from datetime import datetime, timedelta
from typing import Dict

from test_pipi_batch import execute_multi_round_session

# ─── 配置 ─────────────────────────────────────────
STATE_FILE = "cross_session_state.json"
RESULTS_DIR = "timed_results"
DEFAULT_INTERVALS = [1, 24, 72]  # 默认间隔: 1h, 1d, 3d
STATE_VERSION = 1

# ─── 状态管理 ─────────────────────────────────────

def load_state() -> Dict:
    """加载状态文件，不存在返回空字典"""
    if os.path.exists(STATE_FILE):
        with open(STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}

def save_state(state: Dict):
    """保存状态文件"""
    with open(STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2)
    print(f"[状态] 已保存: {STATE_FILE}")

def get_now_iso() -> str:
    """获取当前时间 ISO 格式"""
    return datetime.now().isoformat()

def parse_iso(iso_str: str) -> datetime:
    """兼容 Python 3.6 的 ISO 时间解析"""
    if not iso_str:
        return datetime.min
    try:
        return datetime.fromisoformat(iso_str)
    except AttributeError:
        # Python 3.6 不支持 fromisoformat，手动解析
        try:
            # 处理时区如 +08:00
            s = iso_str
            if '+' in s:
                s = s[:s.index('+')]
            elif s.count('-') > 2:
                # 末尾有 -HH:MM 形式的时区
                last_dash = s.rfind('-')
                if last_dash > 10:
                    s = s[:last_dash]
            return datetime.strptime(s.split('.')[0], "%Y-%m-%dT%H:%M:%S")
        except Exception:
            return datetime.min


# ─── 用例状态结构 ────────────────────────────────

def create_case_state(row: Dict) -> Dict:
    """从 xlsx 的一行创建用例的初始状态。"""
    case_id = row.get("case_id", "UNKNOWN")

    intervals = [
        float(row.get("interval_1", DEFAULT_INTERVALS[0])),
        float(row.get("interval_2", DEFAULT_INTERVALS[1])),
        float(row.get("interval_3", DEFAULT_INTERVALS[2])),
    ]

    session_inputs = {
        "S1": row.get("s1_input", "") or "",
        "S2": row.get("s2_input", "") or "",
        "S3": row.get("s3_input", "") or "",
        "S4": row.get("s4_input", "") or "",
    }

    sessions = []
    for i, sid in enumerate(["S1", "S2", "S3", "S4"]):
        if i == 0:
            status = "pending"
            execute_after = get_now_iso()
        else:
            prev_interval = intervals[i - 1] if (i - 1) < len(intervals) else 0
            if prev_interval > 0:
                status = "blocked"
                execute_after = ""
            else:
                status = "skipped"
                execute_after = ""
        sessions.append({
            "id": sid,
            "status": status,
            "execute_after": execute_after,
            "executed_at": "",
            "response": "",
            "response_time_ms": 0,
            "error": "",
            "retry_count": 0
        })

    return {
        "case_id": case_id,
        "persona": row.get("persona", ""),
        "device_id": row.get("device_id", ""),
        "dimension": row.get("dimension", ""),
        "title": row.get("title", ""),
        "intervals_hours": intervals,
        "session_inputs": session_inputs,
        "expected_output": row.get("expected_output", ""),
        "score_10_desc": row.get("score_10_desc", ""),
        "score_8_desc": row.get("score_8_desc", ""),
        "score_6_desc": row.get("score_6_desc", ""),
        "score_4_desc": row.get("score_4_desc", ""),
        "score_2_desc": row.get("score_2_desc", ""),
        "sessions": sessions,
        "created_at": get_now_iso(),
        "updated_at": get_now_iso(),
        "completed": False
    }


# ─── 执行引擎 ─────────────────────────────────────

def execute_session(case_state: Dict, session_id: str) -> Dict:
    """执行一个 session 的 API 调用。"""
    case_id = case_state["case_id"]
    device_id = case_state["device_id"]
    session_input = case_state["session_inputs"].get(session_id, "")

    if not session_input:
        return {
            "status": "skipped",
            "response": "",
            "response_time_ms": 0,
            "error": "input为空"
        }

    print(f"  [执行] {case_id} {session_id} (device={device_id})...")

    try:
        result = execute_multi_round_session(session_input, device_id=device_id)

        all_responses = []
        total_time = 0
        for r in result["round_results"]:
            all_responses.append(r["pipi_response"])
            total_time += r["response_time_ms"]

        return {
            "status": "done",
            "response": "\n".join(all_responses),
            "response_time_ms": round(total_time, 2),
            "error": ""
        }
    except Exception as e:
        return {
            "status": "error",
            "response": "",
            "response_time_ms": 0,
            "error": str(e)
        }


def update_session_after_execution(case_state: Dict, session_idx: int, result: Dict):
    """执行完成后更新 session 状态，并解锁后续可执行的 session。"""
    sessions = case_state["sessions"]
    intervals = case_state["intervals_hours"]

    sessions[session_idx]["status"] = result["status"]
    sessions[session_idx]["executed_at"] = get_now_iso()
    sessions[session_idx]["response"] = result["response"]
    sessions[session_idx]["response_time_ms"] = result["response_time_ms"]
    sessions[session_idx]["error"] = result["error"]

    if result["status"] != "done":
        case_state["updated_at"] = get_now_iso()
        return

    next_idx = session_idx + 1
    while next_idx < len(sessions):
        next_session = sessions[next_idx]
        interval_hours = intervals[next_idx - 1] if (next_idx - 1) < len(intervals) else 0

        if interval_hours > 0:
            execute_time = datetime.now() + timedelta(hours=interval_hours)
            next_session["status"] = "pending"
            next_session["execute_after"] = execute_time.isoformat()
            print(f"  [计划] {next_session['id']} 将在 {interval_hours}h 后执行 ({execute_time.strftime('%m-%d %H:%M')})")
            break
        else:
            if next_session["status"] != "skipped":
                next_session["status"] = "skipped"
                next_session["execute_after"] = get_now_iso()
                next_session["executed_at"] = get_now_iso()
            next_idx += 1

    all_done = all(s["status"] in ("done", "skipped") for s in sessions)
    case_state["completed"] = all_done
    case_state["updated_at"] = get_now_iso()


# ─── 核心命令 ─────────────────────────────────────

def cmd_init(xlsx_path: str):
    """--init: 从 xlsx 读取用例，创建状态文件，执行所有 S1。"""
    from openpyxl import load_workbook

    if not os.path.exists(xlsx_path):
        print(f"[错误] 文件不存在: {xlsx_path}")
        sys.exit(1)

    print(f"[初始化] 从 {xlsx_path} 加载定时用例...")

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["定时用例"]

    headers = [cell.value for cell in ws[1]]

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(headers, row)))

    state = {
        "version": STATE_VERSION,
        "cases": {},
        "created_at": get_now_iso(),
        "last_run": get_now_iso()
    }

    print(f"[初始化] 共 {len(rows)} 条用例")

    for row_data in rows:
        case_state = create_case_state(row_data)
        case_id = case_state["case_id"]
        state["cases"][case_id] = case_state
        print(f"  [注册] {case_id} ({case_state['title'][:30]})")

    save_state(state)

    print("\n[执行] 开始执行所有 S1...")
    cmd_step()


def cmd_step():
    """--step: 增量推进。遍历状态文件，执行所有"到时间"的 pending session。"""
    state = load_state()
    if not state or "cases" not in state:
        print("[错误] 状态文件为空，请先运行 --init")
        return

    cases = state["cases"]
    now = datetime.now()
    executed_count = 0
    error_count = 0

    print(f"[推进] {now.strftime('%Y-%m-%d %H:%M:%S')} - 检查 {len(cases)} 条用例...")

    for case_id, case_state in cases.items():
        if case_state.get("completed", False):
            continue

        for idx, session in enumerate(case_state["sessions"]):
            if session["status"] != "pending":
                continue

            execute_after = parse_iso(session["execute_after"])
            if now < execute_after:
                continue

            result = execute_session(case_state, session["id"])
            update_session_after_execution(case_state, idx, result)

            if result["status"] == "done":
                executed_count += 1
                print(f"  ✅ {case_id} {session['id']} - {result['response_time_ms']}ms")
            else:
                error_count += 1
                print(f"  ❌ {case_id} {session['id']} - {result['error']}")

    if executed_count == 0 and error_count == 0:
        print("  (没有到执行时间的用例)")

    state["last_run"] = get_now_iso()
    save_state(state)

    total = len(cases)
    completed = sum(1 for c in cases.values() if c.get("completed", False))
    print(f"\n[完成] 本轮: {executed_count} 执行, {error_count} 错误")
    print(f"[进度] {completed}/{total} 用例已完成")
    print(f"[剩余] {total - completed} 用例仍在推进中")

    save_timed_results(state)


def cmd_status():
    """--status: 查看当前所有定时用例的进度。"""
    state = load_state()
    if not state or "cases" not in state:
        print("[状态] 空 — 尚未初始化")
        return

    cases = state["cases"]
    total = len(cases)
    completed = sum(1 for c in cases.values() if c.get("completed", False))

    print(f"\n{'='*60}")
    print(f"定时用例状态总览")
    print(f"创建时间: {state.get('created_at', 'N/A')}")
    print(f"上次运行: {state.get('last_run', 'N/A')}")
    print(f"{'='*60}")
    print(f"总用例: {total} | 已完成: {completed} | 进行中: {total - completed}")
    print(f"{'='*60}\n")

    for case_id, case_state in cases.items():
        sessions_status = []
        for s in case_state["sessions"]:
            icon = "✅" if s["status"] == "done" else "⏳" if s["status"] == "pending" else "🔒" if s["status"] == "blocked" else "❌"
            sessions_status.append(f"{icon}{s['id']}")

        eta = ""
        if not case_state.get("completed", False):
            for s in case_state["sessions"]:
                if s["status"] == "pending" and s["execute_after"]:
                    eta_time = parse_iso(s["execute_after"])
                    if eta_time > datetime.min:
                        remaining = eta_time - datetime.now()
                        if remaining.total_seconds() > 0:
                            hours = int(remaining.total_seconds() / 3600)
                            minutes = int((remaining.total_seconds() % 3600) / 60)
                            eta = f" (剩余 {hours}h{minutes}min)"
                        else:
                            eta = " (已到期，等待执行)"
                    break

        print(f"  {' | '.join(sessions_status)}  {case_id}  {case_state.get('title', '')[:35]} {eta}")


def cmd_retry(case_id: str, session_id: str):
    """--retry: 重跑指定用例的某个 session。"""
    state = load_state()
    if case_id not in state.get("cases", {}):
        print(f"[错误] 未找到用例: {case_id}")
        return

    case_state = state["cases"][case_id]
    for idx, s in enumerate(case_state["sessions"]):
        if s["id"] == session_id:
            s["status"] = "pending"
            s["execute_after"] = get_now_iso()
            s["retry_count"] = s.get("retry_count", 0) + 1
            s["error"] = ""
            print(f"[重试] {case_id} {session_id} (第 {s['retry_count']} 次重试)")
            save_state(state)
            return

    print(f"[错误] 未找到 session: {session_id}")


# ─── 结果导出 ─────────────────────────────────────

def save_timed_results(state: Dict):
    """将当前状态导出为 JSON 结果文件"""
    os.makedirs(RESULTS_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    results = []
    for case_id, case_state in state.get("cases", {}).items():
        case_result = {
            "case_id": case_id,
            "persona": case_state.get("persona", ""),
            "dimension": case_state.get("dimension", ""),
            "title": case_state.get("title", ""),
            "completed": case_state.get("completed", False),
            "sessions": []
        }

        for s in case_state.get("sessions", []):
            if s["status"] in ("done", "error", "skipped"):
                case_result["sessions"].append({
                    "id": s["id"],
                    "status": s["status"],
                    "executed_at": s["executed_at"],
                    "response": s["response"],
                    "response_time_ms": s["response_time_ms"],
                    "error": s["error"]
                })

        results.append(case_result)

    result_file = os.path.join(RESULTS_DIR, f"timed_results_{timestamp}.json")
    with open(result_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"[结果] 已保存: {result_file}")


# ─── 主入口 ─────────────────────────────────────

def print_usage():
    print("""
memory_test_runner.py - 跨会话长期记忆测试定时执行器

用法:
  # 初始化（从 xlsx 创建状态 + 执行所有 S1）
  python3 memory_test_runner.py --init test_cases_timed.xlsx

  # 增量推进（检查状态，执行到期的 session）
  python3 memory_test_runner.py --step

  # 查看当前进度
  python3 memory_test_runner.py --status

  # 强制重跑某个 session
  python3 memory_test_runner.py --retry C2-01 --session S2
""")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print_usage()
        sys.exit(1)

    command = sys.argv[1]

    if command == "--init":
        xlsx_path = sys.argv[2] if len(sys.argv) > 2 else "test_cases_timed.xlsx"
        cmd_init(xlsx_path)
    elif command == "--step":
        cmd_step()
    elif command == "--status":
        cmd_status()
    elif command == "--retry":
        case_id = sys.argv[2] if len(sys.argv) > 2 else ""
        session_id = ""
        if "--session" in sys.argv:
            idx = sys.argv.index("--session")
            session_id = sys.argv[idx + 1] if idx + 1 < len(sys.argv) else ""
        if case_id and session_id:
            cmd_retry(case_id, session_id)
        else:
            print("用法: --retry <case_id> --session <S1|S2|S3|S4>")
    else:
        print_usage()
