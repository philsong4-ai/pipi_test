"""将 7 个用户测试用例（共 770 条）放到单 Sheet Excel"""

import json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('/tmp/pipi_cases.json', 'r', encoding='utf-8') as f:
    cases = json.load(f)

with open('/tmp/pipi_personas.json', 'r', encoding='utf-8') as f:
    personas = json.load(f)

# 把 persona 元信息注入到 cases 里
for c in cases:
    pid = c['persona_id']
    p = personas.get(pid, {})
    c['persona_name'] = p.get('name', '')
    c['persona_created_at'] = p.get('created_at', '')
    c['target_api'] = p.get('target_api', 'pipi')

print(f'loaded {len(cases)} cases across {len(personas)} personas')

wb = Workbook()

# ─── 样式 ────────────────────────────────────
TITLE_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
REDTEAM_FILL = PatternFill("solid", fgColor="FCE4EC")
DIM_FILL = PatternFill("solid", fgColor="E7E6E6")
USER_SEP_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
LEFT_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="left")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ─── 按用户分组 ─────────────────────────────
by_user = defaultdict(list)
for c in cases:
    by_user[c['persona_id']].append(c)

# 按用户创建时间倒序
def user_sort_key(pid):
    cs = by_user[pid]
    return cs[0].get('persona_created_at') or '0'

persona_ids_sorted = sorted(by_user.keys(), key=user_sort_key, reverse=True)

# ─── 单个 Sheet 装所有用例 ──────────────────
ws = wb.active
ws.title = "7用户用例"

ws["A1"] = f"7 个用户测试用例（共 {len(cases)} 条）"
ws["A1"].font = TITLE_FONT
ws["A1"].fill = TITLE_FILL
ws["A1"].alignment = LEFT_CENTER
ws.merge_cells("A1:O1")
ws.row_dimensions[1].height = 28

# 表头
headers = ["序号", "用户ID", "用户名", "创建时间", "target_api",
           "维度", "case_id", "标题", "优先级", "是否红队",
           "红队陷阱类型", "用户输入", "期望输出", "评估点", "不该踩的雷"]
HEADER_ROW = 3
for i, h in enumerate(headers, 1):
    c = ws.cell(row=HEADER_ROW, column=i, value=h)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = BORDER
ws.row_dimensions[HEADER_ROW].height = 22

current_row = HEADER_ROW + 1
global_idx = 0

for user_idx, pid in enumerate(persona_ids_sorted, start=1):
    user_cases = by_user[pid]
    user_cases.sort(key=lambda x: (x.get('dimension_code') or '', x.get('case_id') or ''))
    name = user_cases[0].get('persona_name') or ''
    created = user_cases[0].get('persona_created_at') or ''
    target_api = user_cases[0].get('target_api') or ''

    # 用户分隔行
    ws.cell(row=current_row, column=1,
            value=f"—— 用户 {user_idx}/{len(persona_ids_sorted)}：{name}（{pid}） - 创建时间 {created} - target_api={target_api} - 共 {len(user_cases)} 条 ——").alignment = LEFT_CENTER
    ws.cell(row=current_row, column=1).fill = USER_SEP_FILL
    ws.cell(row=current_row, column=1).font = Font(bold=True, color="7F6000", size=11)
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=15)
    ws.row_dimensions[current_row].height = 24
    current_row += 1

    for c in user_cases:
        global_idx += 1
        is_redteam = c.get('is_redteam') in ('1', '1.0', 1, True)
        ws.cell(row=current_row, column=1, value=global_idx).alignment = CENTER
        ws.cell(row=current_row, column=2, value=pid).alignment = LEFT_CENTER
        ws.cell(row=current_row, column=3, value=name).alignment = LEFT_CENTER
        ws.cell(row=current_row, column=4, value=created).alignment = LEFT_CENTER
        ws.cell(row=current_row, column=5, value=target_api).alignment = CENTER
        ws.cell(row=current_row, column=6, value=c.get('dimension_code') or '').alignment = CENTER
        ws.cell(row=current_row, column=7, value=c.get('case_id') or '').alignment = LEFT_CENTER
        ws.cell(row=current_row, column=8, value=c.get('title') or '').alignment = WRAP
        ws.cell(row=current_row, column=9, value=c.get('priority') or '').alignment = CENTER
        ws.cell(row=current_row, column=10, value="是" if is_redteam else "否").alignment = CENTER
        ws.cell(row=current_row, column=11, value=c.get('redteam_trap_type') or '').alignment = WRAP
        ws.cell(row=current_row, column=12, value=c.get('input_text') or '').alignment = WRAP
        ws.cell(row=current_row, column=13, value=c.get('expected_output') or '').alignment = WRAP
        ws.cell(row=current_row, column=14, value=c.get('evaluation_points') or '').alignment = WRAP
        ws.cell(row=current_row, column=15, value=c.get('failure_flags') or '').alignment = WRAP

        for col in range(1, 16):
            ws.cell(row=current_row, column=col).border = BORDER
            if is_redteam:
                ws.cell(row=current_row, column=col).fill = REDTEAM_FILL
            elif col == 6 and c.get('dimension_code'):
                ws.cell(row=current_row, column=col).fill = DIM_FILL

        ws.row_dimensions[current_row].height = 100
        current_row += 1

# 列宽
widths = [6, 24, 18, 20, 10, 7, 12, 28, 8, 8, 16, 50, 50, 36, 30]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# 冻结表头 + 前两列
ws.freeze_panes = "C4"

output = "/Users/songxuewu/.claude/projects/-Users-songxuewu/pipi-test/7用户测试用例_770条.xlsx"
wb.save(output)
print(f"✓ Excel 已生成: {output}")
print(f"  - 单 Sheet，共 {global_idx} 条用例（{len(persona_ids_sorted)} 个用户）")
print(f"  - 按用户创建时间倒序，每组前有黄色分隔行")
